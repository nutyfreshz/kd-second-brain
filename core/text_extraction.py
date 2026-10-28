from __future__ import annotations

from io import BytesIO, StringIO
from html.parser import HTMLParser
import csv
import json
from pathlib import Path
import re
from typing import Iterable


_TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl",
    ".yaml", ".yml", ".xml", ".html", ".htm", ".sql", ".py", ".r",
    ".js", ".jsx", ".ts", ".tsx", ".java", ".scala", ".sh", ".ps1",
    ".bat", ".log", ".ini", ".cfg", ".conf", ".toml", ".properties",
}


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp874", "windows-1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


class _HTMLText(HTMLParser):
    BLOCKS = {
        "p", "div", "section", "article", "header", "footer", "main", "aside",
        "li", "tr", "table", "br", "h1", "h2", "h3", "h4", "h5", "h6",
    }

    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self.heading_level: int | None = None

    def handle_starttag(self, tag, attrs):
        if tag in self.BLOCKS:
            self.parts.append("\n")
        if re.fullmatch(r"h[1-6]", tag):
            self.heading_level = int(tag[1])
            self.parts.append("#" * self.heading_level + " ")

    def handle_endtag(self, tag):
        if tag in self.BLOCKS:
            self.parts.append("\n")
        if re.fullmatch(r"h[1-6]", tag):
            self.heading_level = None

    def handle_data(self, data):
        text = data.strip()
        if text:
            self.parts.append(text + " ")

    def text(self) -> str:
        raw = "".join(self.parts)
        raw = re.sub(r"[ \t]+\n", "\n", raw)
        raw = re.sub(r"\n{3,}", "\n\n", raw)
        return raw.strip()


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    parts: list[str] = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(f"# Page {i}\n\n{text}")
    return "\n\n".join(parts)


def _paragraph_to_markdown(paragraph) -> str:
    text = paragraph.text.strip()
    if not text:
        return ""
    style = str(getattr(paragraph.style, "name", "") or "")
    m = re.match(r"Heading\s*([1-6])", style, re.I)
    if m:
        return f"{'#' * int(m.group(1))} {text}"
    return text


def _table_to_markdown(rows: Iterable[Iterable[str]]) -> str:
    rows = [[str(cell).replace("\n", " ").strip() for cell in row] for row in rows]
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    header = padded[0]
    out = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    out.extend("| " + " | ".join(row) + " |" for row in padded[1:])
    return "\n".join(out)


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(data))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        rendered = _paragraph_to_markdown(paragraph)
        if rendered:
            parts.append(rendered)
    for i, table in enumerate(doc.tables, 1):
        rendered = _table_to_markdown(
            [[cell.text for cell in row.cells] for row in table.rows]
        )
        if rendered:
            parts.append(f"## Table {i}\n\n{rendered}")
    return "\n\n".join(parts)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(BytesIO(data))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        texts: list[str] = []
        title = ""
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            value = str(getattr(shape, "text", "") or "").strip()
            if not value:
                continue
            if shape == getattr(slide.shapes, "title", None):
                title = value
            else:
                texts.append(value)
        heading = f"# Slide {i}" + (f": {title}" if title else "")
        parts.append(heading + ("\n\n" + "\n".join(texts) if texts else ""))
    return "\n\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=False)
    parts: list[str] = []
    try:
        for ws in wb.worksheets:
            lines = [f"# Sheet: {ws.title}"]
            for row in ws.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    lines.append("\t".join(values))
            parts.append("\n".join(lines))
    finally:
        wb.close()
    return "\n\n".join(parts)


def _extract_delimited(data: bytes, delimiter: str) -> str:
    text = _decode_text(data)
    reader = csv.reader(StringIO(text), delimiter=delimiter)
    return "\n".join("\t".join(cell.strip() for cell in row) for row in reader)


def extract_text_from_bytes(
    *,
    name: str,
    mime_type: str | None,
    data: bytes,
) -> str | None:
    """Best-effort extraction for common text-bearing files.

    Returns None for unsupported binary formats instead of indexing garbage bytes.
    Scanned PDFs/images require a separate OCR/multimodal ingestion path.
    """
    suffix = Path(name).suffix.lower()
    mime = (mime_type or "").lower()

    if suffix == ".pdf" or mime == "application/pdf":
        return _extract_pdf(data)
    if suffix == ".docx" or mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _extract_docx(data)
    if suffix == ".pptx" or mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        return _extract_pptx(data)
    if suffix == ".xlsx" or mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return _extract_xlsx(data)
    if suffix == ".csv" or mime in {"text/csv", "application/csv"}:
        return _extract_delimited(data, ",")
    if suffix == ".tsv" or mime in {"text/tab-separated-values"}:
        return _extract_delimited(data, "\t")
    if suffix in {".html", ".htm"} or mime == "text/html":
        parser = _HTMLText()
        parser.feed(_decode_text(data))
        return parser.text()
    if suffix in {".json", ".jsonl"} or mime in {"application/json", "application/ld+json"}:
        text = _decode_text(data)
        if suffix == ".jsonl":
            return text
        try:
            return json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            return text
    if mime.startswith("text/") or suffix in _TEXT_EXTENSIONS:
        return _decode_text(data).strip()

    # Generic fallback: do not require a known extension when the payload itself
    # is clearly textual. This keeps ingestion format-agnostic while refusing
    # typical binary blobs.
    if data and b"\x00" not in data[:8192]:
        decoded = _decode_text(data).strip()
        if decoded:
            sample = decoded[:12000]
            acceptable = sum(
                1 for ch in sample
                if ch.isprintable() or ch in "\n\r\t"
            )
            replacement_ratio = sample.count("\ufffd") / max(1, len(sample))
            if acceptable / max(1, len(sample)) >= 0.92 and replacement_ratio <= 0.01:
                return decoded

    return None


def extract_text_from_path(path: Path) -> str | None:
    if not path.is_file():
        return None
    return extract_text_from_bytes(
        name=path.name,
        mime_type=None,
        data=path.read_bytes(),
    )


def is_markdown_name(name: str) -> bool:
    return Path(name).suffix.lower() in {".md", ".markdown"}