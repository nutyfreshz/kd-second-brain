import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from core.kb_sync import LocalKnowledgeSource
from core.text_extraction import extract_text_from_bytes


class TextExtractionTests(unittest.TestCase):
    def test_plain_text_and_html(self):
        txt = extract_text_from_bytes(
            name="note.txt",
            mime_type="text/plain",
            data="สวัสดี incentive".encode("utf-8"),
        )
        self.assertIn("incentive", txt)

        html = extract_text_from_bytes(
            name="page.html",
            mime_type="text/html",
            data=b"<h1>Policy</h1><p>Team payout 30</p>",
        )
        self.assertIn("# Policy", html)
        self.assertIn("Team payout 30", html)

    def test_docx_preserves_headings_and_tables(self):
        from docx import Document

        doc = Document()
        doc.add_heading("Incentive Rules", level=1)
        doc.add_paragraph("Team incentive depends on KPI.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "KPI"
        table.cell(0, 1).text = "Threshold"
        table.cell(1, 0).text = "On-time"
        table.cell(1, 1).text = ">98%"
        buf = BytesIO()
        doc.save(buf)

        text = extract_text_from_bytes(
            name="rules.docx",
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data=buf.getvalue(),
        )
        self.assertIn("# Incentive Rules", text)
        self.assertIn("| KPI | Threshold |", text)
        self.assertIn(">98%", text)

    def test_pptx_preserves_slide_boundaries(self):
        from pptx import Presentation

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "MA Incentive"
        slide.placeholders[1].text = "Zone incentive allocation"
        buf = BytesIO()
        prs.save(buf)

        text = extract_text_from_bytes(
            name="deck.pptx",
            mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            data=buf.getvalue(),
        )
        self.assertIn("# Slide 1: MA Incentive", text)
        self.assertIn("Zone incentive allocation", text)

    def test_xlsx_preserves_sheet_names_and_cells(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Install"
        ws.append(["KPI", "Value"])
        ws.append(["On-time", 0.99])
        buf = BytesIO()
        wb.save(buf)

        text = extract_text_from_bytes(
            name="rules.xlsx",
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            data=buf.getvalue(),
        )
        self.assertIn("# Sheet: Install", text)
        self.assertIn("On-time", text)

    def test_local_source_recurses_and_accepts_non_markdown(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            nested = root / "nested"
            nested.mkdir()
            (root / "a.md").write_text("# A\nmarkdown", encoding="utf-8")
            (nested / "b.txt").write_text("plain text rule", encoding="utf-8")
            (nested / "c.json").write_text('{"rule": "json text"}', encoding="utf-8")

            files = LocalKnowledgeSource(td).read_all()
            by_id = {item.source_id: item for item in files}

            self.assertIn("a.md", by_id)
            self.assertIn("nested/b.txt", by_id)
            self.assertIn("nested/c.json", by_id)
            self.assertTrue(by_id["a.md"].use_frontmatter)
            self.assertFalse(by_id["nested/b.txt"].use_frontmatter)

    def test_unknown_extension_is_ingested_when_payload_is_textual(self):
        text = extract_text_from_bytes(
            name="policy.custom",
            mime_type="application/octet-stream",
            data="กฎ incentive แบบข้อความธรรมดา".encode("utf-8"),
        )
        self.assertIn("incentive", text)

    def test_binary_payload_is_not_indexed_as_text(self):
        text = extract_text_from_bytes(
            name="blob.bin",
            mime_type="application/octet-stream",
            data=b"\x00\x01\x02\x03PK\x04\x05",
        )
        self.assertIsNone(text)


if __name__ == "__main__":
    unittest.main()